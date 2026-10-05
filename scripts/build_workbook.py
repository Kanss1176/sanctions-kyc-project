import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

SRC = "powerbi/source/"
OUT = "excel/Citi_Onboarding_Ops_Pack.xlsx"
FOOTER = "Simulated workflow on real GLEIF entity data. Not Citi data."

req = pd.read_csv(SRC + "fact_request.csv")
clock = pd.read_csv(SRC + "sim_clock.csv")["asof_ts"].iloc[0]
for c in req.columns:
    if c.endswith("_ts"):
        req[c] = pd.to_datetime(req[c])
req = req.astype(object).where(req.notna(), None)

wb = Workbook()
ws = wb.active
ws.title = "README"
rows = [
    ("Citi Onboarding Operations Pack", ""),
    ("Label", FOOTER),
    ("Simulation clock", str(clock)),
    ("Entity data", "Real public GLEIF Golden Copy, published 2026-10-04 16:00"),
    ("Requests, errors, SLAs, staffing", "Simulated assumptions (baseline 6 makers, 3 checkers)"),
    ("Version", "v0.1"),
]
for r in rows:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 90

d = wb.create_sheet("Data_Requests")
d.append(list(req.columns))
for r in req.itertuples(index=False):
    d.append(list(r))
for i, c in enumerate(req.columns, 1):
    d.column_dimensions[get_column_letter(i)].width = 20
    if c.endswith("_ts"):
        for cell in d[get_column_letter(i)][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
t = Table(displayName="tblRequests", ref=d.dimensions)
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
d.add_table(t)
d.freeze_panes = "A2"
d.oddFooter.center.text = FOOTER

os.makedirs("excel", exist_ok=True)
wb.save(OUT)
print("saved", OUT, "request rows:", len(req))
