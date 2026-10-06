import json
import shutil

import duckdb
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = "excel/Citi_Onboarding_Ops_Pack.xlsx"
FOOTER = "Simulated workflow on real GLEIF entity data. Not Citi data."
AS_OF = "2026-10-04 16:00:00"
shutil.copy(OUT, "excel/OLD_before_regwatch.xlsx")

fatf = json.load(open("rules/fatf_lists.json", encoding="utf-8"))
con = duckdb.connect("data/processed/onboarding.duckdb", read_only=True)
base = "SELECT COUNT(*) FROM dim_entity WHERE registration_status = 'ISSUED' AND next_renewal_ts < "
v10 = con.sql(base + "TIMESTAMP '" + AS_OF + "'").fetchone()[0]
v11 = con.sql(base + "TIMESTAMP '" + AS_OF + "' - INTERVAL 30 DAY").fetchone()[0]

H = ["change_id", "source", "source_date", "what_changed", "impact", "action", "owner", "status",
     "version_before", "version_after", "evidence"]
rows = [
 ["REG01", "GLEIF Golden Copy", "2026-10-04 16:00", "Entity data baseline recorded (no earlier release compared)",
  "Sets the as-of time for rules R01 to R05", "Logged in provenance register", "Analyst", "Closed", "n/a", "n/a", "Observed"],
 ["REG02", "OFAC, UN and EU sanctions lists", "2026-09-29", "File versions saved as the screening baseline",
  "Used for rule R06 (counts only)", "Hashes and row counts in provenance register", "Analyst", "Closed", "n/a", "n/a", "Observed"],
 ["REG03", "FATF lists", str(fatf["as_of"]), "Country lists saved as the baseline for rule R07",
  "Used for rule R07", "Stored in rules/fatf_lists.json", "Analyst", "Closed", "n/a", "n/a", "Observed"],
 ["REG04", "Simulated internal policy change", "simulated",
  "R02 SLA reduced from 48 hours to 24 hours (change request CR-001)",
  "R02 hit counts unchanged. The SLA is a catalogue value that code does not read. Procedure and escalation matrix updated",
  "Rule table, SOP-R02-01 v1.1 and SOP-ESC-01 v1.1 updated. Tests re-run, 50 passed (simulated record; rule engine unchanged)", "Analyst",
  "Closed (simulated)", "v1.0", "v1.1", "Simulated"],
]

wb = load_workbook(OUT)
if "Reg_Watch" in wb.sheetnames:
    del wb["Reg_Watch"]
w = wb.create_sheet("Reg_Watch")
w["A1"] = "Regulatory and procedure change register"
w["A1"].font = Font(bold=True, size=14)
w["A2"] = "Observed rows record real source dates. REG04 is a simulated change; the rule engine itself was not changed."
for j, h in enumerate(H, 1):
    w.cell(row=4, column=j, value=h)
for i, r in enumerate(rows, 5):
    for j, v in enumerate(r, 1):
        w.cell(row=i, column=j, value=v)
last = 4 + len(rows)
t = Table(displayName="RegWatch", ref="A4:K%d" % last)
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
w.add_table(t)
w.cell(row=last + 2, column=1, value="Blank status or version cells")
w.cell(row=last + 2, column=2, value="=COUNTBLANK(H5:J%d)" % last)
w.cell(row=last + 2, column=3, value='=IF(B%d=0,"OK","CHECK")' % (last + 2))
for c, wd in zip("ABCDEFGHIJK", (10, 30, 18, 52, 46, 52, 10, 18, 14, 14, 12)):
    w.column_dimensions[c].width = wd
w.oddFooter.center.text = FOOTER
wb.save(OUT)
print("Reg_Watch built. FATF as_of:", fatf["as_of"])
