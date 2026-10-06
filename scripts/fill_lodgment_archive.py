import ast
import datetime as dt

import duckdb
from openpyxl import load_workbook

SRC = open("scripts/add_lodgment_archive.py", encoding="utf-8").read()
CONST = {}
for node in ast.parse(SRC).body:
    if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
        name = node.targets[0].id
        if name in ("OUT", "N", "DOCS", "MAKERS", "CHECKERS"):
            try:
                CONST[name] = ast.literal_eval(node.value)
            except Exception:
                print("cannot read", name, "->", ast.unparse(node.value))
print(CONST)
missing = [k for k in ("OUT", "N", "DOCS") if k not in CONST]
assert not missing, "send me this output; missing: " + str(missing)


def as_list(v):
    return v.split(",") if isinstance(v, str) else list(v)


OUT, N = CONST["OUT"], CONST["N"]
DOCS = as_list(CONST["DOCS"])
MAKERS = ["M%d" % i for i in range(1, 9)]
CHECKERS = ["C%d" % i for i in range(1, 5)]
con = duckdb.connect("data/processed/onboarding.duckdb", read_only=True)
IDS = [r[0] for r in con.sql("SELECT request_id FROM fact_request WHERE status = 'Completed' ORDER BY request_id").fetchall()]
IDS = [IDS[i * 100] for i in range(18)]

days = [dt.date(2026, 10, d) for d in range(5, 31) if dt.date(2026, 10, d).weekday() < 5]
assert N >= 18, "N is too small"
wb = load_workbook(OUT)
L, A = wb["Lodgment_Log"], wb["Archive_Manifest"]
ar = 2
for i in range(18):
    maker, checker = MAKERS[i % len(MAKERS)], CHECKERS[i % len(CHECKERS)]
    doc = DOCS[i % len(DOCS)]
    check, status, note = "Complete", "Lodged", None
    if i == 12:
        check, status, note = "Incomplete", "Returned to client", "Returned: required document missing (simulated)."
    if i == 13:
        check, status, note = "Rejected", "Returned to client", "Returned: signatory authority not evidenced (simulated)."
    if i == 14:
        status, note = "Pending", "Awaiting client document (simulated)."
    if i == 15:
        check, note = "Incomplete", "Deliberate test case: lodged without a complete check; the flag should fire."
    received = dt.datetime.combine(days[i], dt.time(9 + i % 6, 0))
    lodged = received + dt.timedelta(hours=2 + i % 5)
    r = i + 2
    L.cell(row=r, column=2, value=IDS[i])
    L.cell(row=r, column=3, value=doc)
    L.cell(row=r, column=4, value=received)
    L.cell(row=r, column=5, value=checker)
    L.cell(row=r, column=6, value=check)
    L.cell(row=r, column=7, value=maker)
    if status == "Lodged":
        L.cell(row=r, column=8, value=lodged)
    L.cell(row=r, column=9, value=status)
    if note:
        L.cell(row=r, column=12, value=note)
    if i < 12 or i in (16, 17):
        exp = 2 + i % 2
        pres = exp if i < 12 else exp - 1
        A.cell(row=ar, column=2, value=IDS[i])
        A.cell(row=ar, column=3, value=doc)
        A.cell(row=ar, column=4, value=lodged.date())
        A.cell(row=ar, column=5, value=2 if i == 5 else 1)
        A.cell(row=ar, column=7, value=exp)
        A.cell(row=ar, column=8, value=pres)
        if i != 16:
            A.cell(row=ar, column=11, value=maker)
            A.cell(row=ar, column=12, value=lodged.date() + dt.timedelta(days=1))
        ar += 1
out2 = OUT.replace(".xlsx", "_filled.xlsx")
wb.save(out2)
print("saved", out2, "| lodgment rows 18 | archive rows", ar - 2)
