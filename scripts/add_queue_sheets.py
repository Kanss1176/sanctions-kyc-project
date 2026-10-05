import duckdb
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "excel/Citi_Onboarding_Ops_Pack.xlsx"
FOOTER = "Simulated workflow on real GLEIF entity data. Not Citi data."
con = duckdb.connect("data/processed/onboarding.duckdb", read_only=True)
clock = con.sql("SELECT asof_ts FROM sim_clock").fetchone()[0]
openq = con.sql("SELECT request_id, request_type, received_ts, due_ts, maker_id FROM fact_request WHERE status='Open' ORDER BY due_ts").fetchall()
errs = con.sql("SELECT error_code, description FROM dim_error_type ORDER BY 1").fetchall()

wb = load_workbook(OUT)
for n in ("Daily_Queue", "Checker_Log", "Lists"):
    if n in wb.sheetnames:
        del wb[n]
wb["Data_Requests"].column_dimensions["B"].width = 26
bold = Font(bold=True)

L = wb.create_sheet("Lists")
L.append(["error_code", "description"])
for e in errs:
    L.append(list(e))
n_err = len(errs) + 1

q = wb.create_sheet("Daily_Queue")
q["A1"], q["B1"] = "Simulation clock", clock
q["B1"].number_format = "yyyy-mm-dd hh:mm"
q["C1"] = "Hours left = calendar hours to due time. Priority bands are my assumption."
heads = ["request_id", "request_type", "received_ts", "due_ts", "maker_id", "hours_left", "priority_band", "action_status"]
for j, h in enumerate(heads, 1):
    q.cell(row=3, column=j, value=h).font = bold
for i, r in enumerate(openq, 4):
    for j, v in enumerate(r, 1):
        q.cell(row=i, column=j, value=v)
    q.cell(row=i, column=3).number_format = "yyyy-mm-dd hh:mm"
    q.cell(row=i, column=4).number_format = "yyyy-mm-dd hh:mm"
    q.cell(row=i, column=6, value="=(D%d-$B$1)*24" % i).number_format = "0.0"
    q.cell(row=i, column=7, value='=IF(F%d<0,"P0 Overdue",IF(F%d<24,"P1 under 24h",IF(F%d<72,"P2 under 72h","P3 later")))' % (i, i, i))
    q.cell(row=i, column=8, value="Not started")
last = 3 + len(openq)
rng = "F4:F%d" % last
q.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", bgColor="FF9999")))
q.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["24"], fill=PatternFill("solid", bgColor="FFE699")))
dv = DataValidation(type="list", formula1='"Not started,In progress,Waiting on client,Done"', allow_blank=True)
q.add_data_validation(dv)
dv.add("H4:H%d" % last)
q.auto_filter.ref = "A3:H%d" % last
q.freeze_panes = "A4"
for c, w in zip("ABCDEFGH", (14, 18, 18, 18, 10, 11, 16, 18)):
    q.column_dimensions[c].width = w
q.oddFooter.center.text = FOOTER

k = wb.create_sheet("Checker_Log")
k.append(["log_id", "request_id", "maker_id", "checker_id", "result", "error_code", "checked_ts", "comment"])
for c in k[1]:
    c.font = bold
for i in range(2, 202):
    rid = openq[i - 2][0] if i - 2 < len(openq) else None
    k.append(["LOG%04d" % (i - 1), rid, '=IF(B%d="","",IFERROR(INDEX(Data_Requests!$F$2:$F$2001,MATCH(B%d,Data_Requests!$A$2:$A$2001,0)),""))' % (i, i)])
v1 = DataValidation(type="custom", formula1='=AND(D2<>"",D2<>C2,LEFT(D2,1)="C")', allow_blank=True, showErrorMessage=True,
                    errorTitle="Maker-checker rule", error="Checker must be a checker ID (C1 to C3) and cannot be the maker.")
v2 = DataValidation(type="list", formula1='"Pass,Rework"', allow_blank=True)
v3 = DataValidation(type="list", formula1="=Lists!$A$2:$A$%d" % n_err, allow_blank=True)
for d, col in ((v1, "D"), (v2, "E"), (v3, "F")):
    k.add_data_validation(d)
    d.add("%s2:%s201" % (col, col))
for c in "GH":
    k.column_dimensions[c].width = 20
for c in "ABCDEF":
    k.column_dimensions[c].width = 14
for r in range(2, 202):
    k.cell(row=r, column=7).number_format = "yyyy-mm-dd hh:mm"
k.freeze_panes = "A2"
k.oddFooter.center.text = FOOTER

wb.save(OUT)
print("added Daily_Queue rows:", len(openq), "Checker_Log rows: 200")
