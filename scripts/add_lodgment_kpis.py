from openpyxl import load_workbook
from openpyxl.styles import Font

OUT = "excel/Citi_Onboarding_Ops_Pack.xlsx"
wb = load_workbook(OUT)
k = wb["KPI_Summary"]
r0 = 34
while k.cell(row=r0, column=1).value is not None:
    r0 += 1
bold = Font(bold=True)
k.cell(row=r0, column=1, value="Lodgment and archive KPIs (simulated logs, Excel only)").font = bold
rows = [
 ("Lodgment entries", "=COUNTA(Lodgment_Log!B2:B201)", "0"),
 ("Lodged", '=COUNTIF(Lodgment_Log!I2:I201,"Lodged")', "0"),
 ("Lodged with a Complete check", '=COUNTIFS(Lodgment_Log!I2:I201,"Lodged",Lodgment_Log!F2:F201,"Complete")', "0"),
 ("Average hours to lodge", "=AVERAGE(Lodgment_Log!J2:J201)", "0.00"),
 ("Lodged within 4 hours (assumed target)", '=COUNTIFS(Lodgment_Log!I2:I201,"Lodged",Lodgment_Log!J2:J201,"<=4")/B{lodged}', "0.0%"),
 ("Lodgment rows with a CHECK flag", '=COUNTIF(Lodgment_Log!K2:K201,"CHECK*")', "0"),
 ("Archive entries", "=COUNTA(Archive_Manifest!B2:B201)", "0"),
 ("Archive complete", '=COUNTIF(Archive_Manifest!I2:I201,"Complete")', "0"),
 ("Archive completeness rate", "=B{acomp}/B{aent}", "0.0%"),
 ("Archive rows with a CHECK flag", '=COUNTIF(Archive_Manifest!M2:M201,"CHECK*")', "0"),
 ("Checker_Log rows with a CHECK flag", '=COUNTIF(Checker_Log!I2:I201,"CHECK*")', "0"),
]
ref = {"lodged": r0 + 2, "aent": r0 + 7, "acomp": r0 + 8}
for i, (name, f, fmt) in enumerate(rows, 1):
    k.cell(row=r0 + i, column=1, value=name)
    c = k.cell(row=r0 + i, column=2, value=f.format(**ref))
    c.number_format = fmt
    k.cell(row=r0 + i, column=3, value="No DuckDB equivalent")
wb.save(OUT)
print("added at row", r0, "to", r0 + len(rows))
