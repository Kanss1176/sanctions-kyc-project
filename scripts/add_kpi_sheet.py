import duckdb
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule

OUT = "excel/Citi_Onboarding_Ops_Pack.xlsx"
FOOTER = "Simulated workflow on real GLEIF entity data. Not Citi data."
con = duckdb.connect("data/processed/onboarding.duckdb", read_only=True)
one = lambda s: con.sql(s).fetchone()[0]
R = lambda c: "Data_Requests!$%s$2:$%s$2001" % (c, c)
bold = Font(bold=True)

M = [
 ("Requests received", "=COUNTA(%s)" % R("A"), one("SELECT COUNT(*) FROM fact_request"), "0"),
 ("Completed", '=COUNTIFS(%s,"Completed")' % R("J"), one("SELECT COUNT(*) FROM fact_request WHERE status='Completed'"), "0"),
 ("Open", '=COUNTIFS(%s,"Open")' % R("J"), one("SELECT COUNT(*) FROM fact_request WHERE status='Open'"), "0"),
 ("Median turnaround (hours)", "=MEDIAN(%s)" % R("K"), one("SELECT quantile_cont(tat_hours,0.5) FROM fact_request"), "0.00"),
 ("P90 turnaround (hours)", "=PERCENTILE(%s,0.9)" % R("K"), one("SELECT quantile_cont(tat_hours,0.9) FROM fact_request"), "0.00"),
 ("SLA breaches (count)", "=COUNTIFS(%s,TRUE)" % R("M"), one("SELECT COUNT(*) FROM fact_request WHERE sla_breached"), "0"),
 ("SLA breach % of completed", '=COUNTIFS(%s,TRUE)/COUNTIFS(%s,"Completed")' % (R("M"), R("J")),
  one("SELECT COUNT(*) FILTER (WHERE sla_breached)*1.0/COUNT(*) FILTER (WHERE status='Completed') FROM fact_request"), "0.0%"),
 ("First-time-right % of completed", '=COUNTIFS(%s,TRUE)/COUNTIFS(%s,"Completed")' % (R("O"), R("J")),
  one("SELECT COUNT(*) FILTER (WHERE first_time_right)*1.0/COUNT(*) FILTER (WHERE status='Completed') FROM fact_request"), "0.0%"),
 ("Requests with an error", "=COUNTIFS(%s,TRUE)" % R("P"), one("SELECT COUNT(*) FROM fact_request WHERE has_error"), "0"),
 ("Error rate (simulated, assumed injected rate)", "=COUNTIFS(%s,TRUE)/COUNTA(%s)" % (R("P"), R("A")),
  one("SELECT COUNT(*) FILTER (WHERE has_error)*1.0/COUNT(*) FROM fact_request"), "0.0%"),
 ("Escalated requests (matches fact_escalation)", "=COUNTIFS(%s,TRUE)" % R("Q"), one("SELECT COUNT(DISTINCT request_id) FROM fact_escalation"), "0"),
 ("Total rework loops", "=SUM(%s)" % R("N"), one("SELECT SUM(rework_count) FROM fact_request"), "0"),
]

wb = load_workbook(OUT)
if "KPI_Summary" in wb.sheetnames:
    del wb["KPI_Summary"]
k = wb.create_sheet("KPI_Summary")
k["A1"] = "KPI summary"
k["A1"].font = Font(bold=True, size=14)
k["A2"] = "Overall check"
k["A2"].font = bold
k["A3"] = "All values are simulated. Escalation and error rates follow my severity and error-injection assumptions."
for j, h in enumerate(["Metric", "Excel formula", "DuckDB value", "Check"], 1):
    k.cell(row=5, column=j, value=h).font = bold
r = 6
for name, f, v, fmt in M:
    k.cell(row=r, column=1, value=name)
    k.cell(row=r, column=2, value=f).number_format = fmt
    k.cell(row=r, column=3, value=float(v) if v is not None else None).number_format = fmt
    k.cell(row=r, column=4, value='=IF(ABS(B%d-C%d)<0.0001,"OK","CHECK")' % (r, r))
    r += 1
m_end = r - 1

r += 1
k.cell(row=r, column=1, value="By request type").font = bold
r += 1
for j, h in enumerate(["request_type", "Count (Excel)", "Count (DuckDB)", "Avg TAT h (Excel)", "Avg TAT h (DuckDB)", "Breach % (Excel)", "Check"], 1):
    k.cell(row=r, column=j, value=h).font = bold
t_start = r + 1
for (t,) in con.sql("SELECT request_type FROM dim_request_type ORDER BY 1").fetchall():
    r += 1
    c, cnt, avg = one("SELECT '%s'" % t), one("SELECT COUNT(*) FROM fact_request WHERE request_type='%s'" % t), one("SELECT AVG(tat_hours) FROM fact_request WHERE request_type='%s'" % t)
    k.cell(row=r, column=1, value=t)
    k.cell(row=r, column=2, value='=COUNTIFS(%s,A%d)' % (R("C"), r))
    k.cell(row=r, column=3, value=cnt)
    k.cell(row=r, column=4, value='=AVERAGEIFS(%s,%s,A%d)' % (R("K"), R("C"), r)).number_format = "0.00"
    k.cell(row=r, column=5, value=float(avg)).number_format = "0.00"
    k.cell(row=r, column=6, value='=COUNTIFS(%s,A%d,%s,TRUE)/COUNTIFS(%s,A%d,%s,"Completed")' % (R("C"), r, R("M"), R("C"), r, R("J"))).number_format = "0.0%"
    k.cell(row=r, column=7, value='=IF(AND(B%d=C%d,ABS(D%d-E%d)<0.001),"OK","CHECK")' % (r, r, r, r))
t_end = r

r += 2
k.cell(row=r, column=1, value="Open requests by priority band (Daily_Queue)").font = bold
r += 1
a_start = r
for band in ("P0*", "P1*", "P2*", "P3*"):
    k.cell(row=r, column=1, value=band.replace("*", ""))
    k.cell(row=r, column=2, value='=COUNTIF(Daily_Queue!$G$4:$G$23,"%s")' % band)
    r += 1
k.cell(row=r, column=1, value="Total in bands")
k.cell(row=r, column=2, value="=SUM(B%d:B%d)" % (a_start, r - 1))
k.cell(row=r, column=3, value="=B8")
k.cell(row=r, column=4, value='=IF(B%d=C%d,"OK","CHECK")' % (r, r))
a_end = r

k["B2"] = '=IF(COUNTIF(D6:D%d,"CHECK")+COUNTIF(G%d:G%d,"CHECK")+COUNTIF(D%d,"CHECK")=0,"ALL OK","REVIEW")' % (m_end, t_start, t_end, a_end)
k["B2"].font = bold
red = PatternFill("solid", bgColor="FF9999")
green = PatternFill("solid", bgColor="C6EFCE")
for rg in ("D6:D%d" % a_end, "G%d:G%d" % (t_start, t_end), "B2"):
    k.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"CHECK"'], fill=red))
    k.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"REVIEW"'], fill=red))
    k.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"OK"'], fill=green))
    k.conditional_formatting.add(rg, CellIsRule(operator="equal", formula=['"ALL OK"'], fill=green))
for c, w in zip("ABCDEFG", (48, 16, 16, 18, 18, 16, 10)):
    k.column_dimensions[c].width = w
k.oddFooter.center.text = FOOTER
wb.save(OUT)
print("KPI_Summary built, metric rows:", len(M), "type rows:", t_end - t_start + 1)
