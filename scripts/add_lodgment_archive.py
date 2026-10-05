from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = "excel/Citi_Onboarding_Ops_Pack.xlsx"
FOOTER = "Simulated workflow on real GLEIF entity data. Not Citi data."
N = 200
bold = Font(bold=True)
red = PatternFill("solid", bgColor="FF9999")
DOCS = "Signatory list,Board resolution,ID document,Account opening form,Signature card,Address proof"
MAKERS = ",".join("M%d" % i for i in range(1, 9))
CHECKERS = ",".join("C%d" % i for i in range(1, 5))

wb = load_workbook(OUT)
for n in ("Lodgment_Log", "Archive_Manifest"):
    if n in wb.sheetnames:
        del wb[n]

def dv_list(ws, rng, items, title):
    d = DataValidation(type="list", formula1=items, allow_blank=True, showErrorMessage=True, errorTitle=title, error="Pick a value from the list.")
    ws.add_data_validation(d)
    d.add(rng)

L = wb.create_sheet("Lodgment_Log")
hdr = ["lodgment_id", "request_id", "document_type", "received_ts", "checked_by", "check_result", "lodged_by", "lodged_ts", "system_status", "hours_to_lodge", "flag", "comment"]
for j, h in enumerate(hdr, 1):
    L.cell(row=1, column=j, value=h).font = bold
for i in range(2, N + 2):
    L.cell(row=i, column=1, value="LDG%04d" % (i - 1))
    L.cell(row=i, column=4).number_format = "yyyy-mm-dd hh:mm"
    L.cell(row=i, column=8).number_format = "yyyy-mm-dd hh:mm"
    L.cell(row=i, column=10, value='=IF(AND(D%d<>"",H%d<>""),ROUND((H%d-D%d)*24,1),"")' % (i, i, i, i))
    L.cell(row=i, column=11, value='=IF(AND(I%d="Lodged",F%d<>"Complete"),"CHECK: lodged without complete check","")' % (i, i))
end = N + 1
dv_list(L, "B2:B%d" % end, "=Data_Requests!$A$2:$A$2001", "request_id")
dv_list(L, "C2:C%d" % end, '"%s"' % DOCS, "document_type")
dv_list(L, "E2:E%d" % end, '"%s"' % CHECKERS, "checker")
dv_list(L, "F2:F%d" % end, '"Complete,Incomplete,Rejected"', "check_result")
dv_list(L, "G2:G%d" % end, '"%s"' % MAKERS, "maker")
dv_list(L, "I2:I%d" % end, '"Pending,Lodged,Returned to client"', "status")
L["N1"] = "Lodgment status"
L["N1"].font = bold
for r, s in enumerate(("Pending", "Lodged", "Returned to client"), 2):
    L.cell(row=r, column=14, value=s)
    L.cell(row=r, column=15, value='=COUNTIF($I$2:$I$%d,N%d)' % (end, r))
L["N6"] = "Rows with CHECK flag"
L["O6"] = '=COUNTIF($K$2:$K$%d,"CHECK*")' % end
L["N8"] = "Simulated log. Checker IDs C1-C4 and maker IDs M1-M8 are assumptions."
L.conditional_formatting.add("K2:K%d" % end, FormulaRule(formula=['LEFT(K2,5)="CHECK"'], fill=red))
for c, w in zip("ABCDEFGHIJKLMNO", (12, 12, 22, 18, 12, 14, 12, 18, 18, 14, 36, 30, 3, 22, 10)):
    L.column_dimensions[c].width = w
L.freeze_panes = "A2"
L.oddFooter.center.text = FOOTER

A = wb.create_sheet("Archive_Manifest")
hdr = ["archive_id", "request_id", "document_type", "doc_date", "version", "file_name", "docs_expected", "docs_present", "completeness", "retention_note", "archived_by", "archived_date", "flag"]
for j, h in enumerate(hdr, 1):
    A.cell(row=1, column=j, value=h).font = bold
for i in range(2, N + 2):
    A.cell(row=i, column=1, value="ARC%04d" % (i - 1))
    A.cell(row=i, column=4).number_format = "yyyy-mm-dd"
    A.cell(row=i, column=12).number_format = "yyyy-mm-dd"
    A.cell(row=i, column=6, value='=IF(OR(B%d="",C%d="",D%d="",E%d=""),"",B%d&"_"&SUBSTITUTE(C%d," ","-")&"_"&TEXT(D%d,"yyyymmdd")&"_v"&E%d)' % ((i,) * 8))
    A.cell(row=i, column=9, value='=IF(OR(G%d="",H%d=""),"",IF(H%d>=G%d,"Complete","Incomplete"))' % ((i,) * 4))
    A.cell(row=i, column=10, value='=IF(B%d="","","Retention period per bank policy; not defined in this simulation")' % i)
    A.cell(row=i, column=13, value='=IF(AND(K%d<>"",I%d<>"Complete"),"CHECK: archived while incomplete","")' % (i, i))
dv_list(A, "B2:B%d" % end, "=Data_Requests!$A$2:$A$2001", "request_id")
dv_list(A, "C2:C%d" % end, '"%s"' % DOCS, "document_type")
dv_list(A, "K2:K%d" % end, '"%s"' % MAKERS, "archived_by")
w = DataValidation(type="whole", operator="between", formula1="0", formula2="20", allow_blank=True, showErrorMessage=True, errorTitle="Count", error="Enter a whole number from 0 to 20.")
A.add_data_validation(w)
w.add("G2:H%d" % end)
A["O1"] = "File naming convention"
A["O1"].font = bold
A["O2"] = "request_id_document-type_YYYYMMDD_vN"
A["O3"] = "Example: REQ01971_Signatory-list_20261030_v1"
A["O4"] = "Archive only when completeness = Complete."
A["O5"] = "Retention period: set by bank policy, not defined here."
A["O6"] = "Simulated manifest. Not Citi data."
A.conditional_formatting.add("M2:M%d" % end, FormulaRule(formula=['LEFT(M2,5)="CHECK"'], fill=red))
A.conditional_formatting.add("I2:I%d" % end, FormulaRule(formula=['I2="Incomplete"'], fill=red))
for c, w_ in zip("ABCDEFGHIJKLMNO", (12, 12, 22, 12, 9, 46, 14, 14, 14, 52, 12, 14, 34, 3, 52)):
    A.column_dimensions[c].width = w_
A.freeze_panes = "A2"
A.oddFooter.center.text = FOOTER

wb.save(OUT)
print("Lodgment_Log and Archive_Manifest built, rows each:", N)
